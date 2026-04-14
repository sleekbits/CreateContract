from __future__ import annotations

from pathlib import Path

from docx import Document


def ensure_sample_template() -> str:
    path = Path("assets/templates/sample_contract_template.docx")
    if path.exists():
        return str(path)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("{{contract_title}}", 0)
    doc.add_paragraph("Contract No: {{contract_no}}")
    doc.add_paragraph("Contract Date: {{contract_date}}")

    table = doc.add_table(rows=6, cols=2)
    rows = [
        ("Client Name", "{{client_name}}"),
        ("Client Address", "{{client_address}}"),
        ("Contractor Name", "{{contractor_name}}"),
        ("Contractor Address", "{{contractor_address}}"),
        ("Start Date", "{{start_date}}"),
        ("End Date", "{{end_date}}"),
    ]
    for i, (k, v) in enumerate(rows):
        table.cell(i, 0).text = k
        table.cell(i, 1).text = v

    doc.add_heading("Commercial Terms", level=2)
    doc.add_paragraph("Contract Value: {{contract_value}}")
    doc.add_paragraph("Payment Terms:")
    doc.add_paragraph("{{payment_terms}}")

    doc.add_heading("Scope of Work", level=2)
    doc.add_paragraph("{{scope_of_work}}")

    doc.add_heading("Special Conditions", level=2)
    doc.add_paragraph("{{special_conditions}}")

    doc.add_paragraph("Prepared By: {{prepared_by}}")
    doc.add_paragraph("Approved By: {{approved_by}}")

    doc.save(path)
    return str(path)
